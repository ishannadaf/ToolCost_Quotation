from tkinter import *
import tkinter.ttk as ttk
from Database.connection import *
from tkinter import messagebox
from openpyxl import Workbook
import os
from openpyxl.styles import Font, Alignment, Border, Side
from collections import defaultdict

def Frm_Delivery_Master(master, login_id):
    class QtyDialog(Toplevel):
        def __init__(self, parent, part_name, available_qty):
            super().__init__(parent)

            self.title("Enter Quantity")
            self.geometry("350x250+400+250")
            self.resizable(False, False)
            self.qty = None          # value to return
            self.available_qty = available_qty

            # Make dialog modal (block interaction with parent until closed)
            self.transient(parent)
            self.grab_set()

            # --------- UI --------- #
            Label(self, text="Delivery Quantity",
                font=("Times New Roman", 16, "bold")).pack(pady=(10, 5))

            Label(self, text=f"Part: {part_name}",
                font=("Times New Roman", 13)).pack(pady=2)

            Label(self, text=f"Available Qty: {available_qty}",
                font=("Times New Roman", 12)).pack(pady=2)

            Label(self, text="Enter Qty to Deliver:",
                font=("Times New Roman", 12)).pack(pady=(10, 3))

            self.entry = Entry(self, font=("Times New Roman", 14),
                            justify="center")
            self.entry.pack(pady=5)
            self.entry.focus()

            # Buttons
            btn_frame = Frame(self)
            btn_frame.pack(pady=15)

            Button(
                btn_frame, text="OK", bg="green", fg="white",
                font=("Times New Roman", 12, "bold"), width=10,
                command=self.on_ok
            ).grid(row=0, column=0, padx=10)

            Button(
                btn_frame, text="Cancel", bg="red", fg="white",
                font=("Times New Roman", 12, "bold"), width=10,
                command=self.on_cancel
            ).grid(row=0, column=1, padx=10)

            # Close dialog on ESC
            self.bind("<Escape>", lambda e: self.on_cancel())

        def on_ok(self):
            text = self.entry.get().strip()
            if not text:
                messagebox.showerror("Invalid Input", "Please enter a quantity.",parent=frm_delivery)
                return
            try:
                value = int(text)
            except ValueError:
                messagebox.showerror("Invalid Input", "Enter a valid integer.",parent=frm_delivery)
                return

            if 1 <= value <= self.available_qty:
                self.qty = value
                self.destroy()
            else:
                messagebox.showerror(
                    "Invalid Range",
                    f"Enter a value between 1 and {self.available_qty}", parent=frm_delivery
                )

        def on_cancel(self):
            self.qty = None
            self.destroy()

    def On_Start():
        TxtQuot.delete(0, END)
        List_Total_TreeView.delete(*List_Total_TreeView.get_children())
        List_Delivery_TreeView.delete(*List_Delivery_TreeView.get_children())
        BtnSave['state'] = 'disabled'

    def Save_Records():
        quot_no = TxtQuot.get()
        all_items_data = []
        chalan_id = Generate_Next_Challan(quot_no)
        # Get the IDs of all top-level items
        top_level_items = List_Delivery_TreeView.get_children()
        dt = datetime.today().date()
        part_list = []
        for item_id in top_level_items:
            # Retrieve information for each item
            item_info = List_Delivery_TreeView.item(item_id)
            values = item_info['values']
            part_list.append(str(values[0]))
            all_items_data.append((chalan_id, quot_no, str(values[0]), str(values[2]), dt, login_id))
        
        if all_items_data:
            sql1 = f"INSERT INTO delivery_manage_master (chalan_id, quot_no, part_no, delivery_qty, delivery_dt, login_id) VALUES ({DATABASE_SYN}, {DATABASE_SYN}, {DATABASE_SYN}, {DATABASE_SYN}, {DATABASE_SYN}, {DATABASE_SYN})"
            db_cursor.executemany(sql1, tuple(all_items_data))
            placeholders = ",".join(f"'{p}'" for p in part_list)
            sql2 = f"""
                    UPDATE quotation_master_details t1 
                inner join quotation_master t3 on t3.quotation_Id = t1.Quotation_Id
                JOIN (
                SELECT quot_no, part_no, SUM(delivery_qty) AS total_qty
                FROM delivery_manage_master
                GROUP BY quot_no, part_no
                ) t2 ON t1.part_no = t2.part_no and t2.quot_no = t3.invoice_id
                SET t1.delivery_flag = '1'
                WHERE t3.invoice_id = '{quot_no}'
                AND t1.part_no IN ({placeholders})
                AND t1.part_qty = t2.total_qty;
            """#
            #print(sql2)
            db_cursor.execute(sql2)
            
            db_connection.commit()
            
            messagebox.showinfo('Success', "Data Saved Successfully",parent=frm_delivery)
            On_Start()

    def Search_Record_Report():
        n1 = TxtInvoiceReport.get()
        if n1:
            sql1 = f"""
                SELECT 
                dmm.chalan_id,
                qmd.part_no,
                qmd.part_desc,
                qmd.part_qty AS Total_Qty,
                dmm.delivery_qty,
                qmd.part_qty 
                - SUM(dmm.delivery_qty) OVER (
                        PARTITION BY dmm.quot_no, dmm.part_no
                        ORDER BY dmm.delivery_dt, dmm.id   -- use your PK instead of dmm.id
                    ) AS remaining_qty,
                dmm.delivery_dt
            FROM quotation_master_details qmd
            inner join quotation_master qm ON qm.quotation_Id = qmd.Quotation_Id
            LEFT JOIN delivery_manage_master dmm
            ON dmm.quot_no = qm.invoice_id
            #INNER JOIN quotation_master qm ON qm.quotation_Id = qmd.Quotation_Id
            AND dmm.part_no = qmd.part_no
            WHERE qm.invoice_id = '{n1}'
            ORDER BY dmm.part_no, dmm.delivery_dt, dmm.id;
            """
            db_cursor.execute(sql1)
            data1 = db_cursor.fetchall()
            
            List_Total_TreeView_report.delete(*List_Total_TreeView_report.get_children())
            if data1:
                for i in data1:
                    List_Total_TreeView_report.insert("", "end", values=(str(i[0]), i[1], str(i[2]), str(i[3]), str(i[4]) if i[4] else '0', str(i[5]) if i[5] else '0', str(i[6]) if i[6] else 'N/A'))
            else:
                messagebox.showinfo("Info", f"No data found for Quotation No. {n1}", parent=frm_delivery)
                
    def Search_Records():
        n1 = TxtQuot.get()
        List_Total_TreeView.delete(*List_Total_TreeView.get_children())
        if n1:
            sql1 = f"""SELECT 
                        qmd.part_no,
                        qmd.part_desc,
                        (qmd.part_qty - COALESCE(SUM(dmm.delivery_qty), 0)) AS remaining_qty
                    FROM quotation_master_details qmd
                    inner join quotation_master qm on qm.quotation_Id = qmd.Quotation_Id
                    LEFT JOIN delivery_manage_master dmm
                    ON dmm.part_no = qmd.part_no
                    AND dmm.quot_no = qm.invoice_id
                    WHERE qm.invoice_id = '{n1}'
                    GROUP BY
                        qmd.part_no,
                        qmd.part_desc,
                        qmd.part_qty
                    HAVING remaining_qty > 0;"""
            db_cursor.execute(sql1)
            data1 = db_cursor.fetchall()
            
            List_Delivery_TreeView.delete(*List_Delivery_TreeView.get_children())
            if data1:
                for i in data1:
                    List_Total_TreeView.insert("", "end", values=(str(i[0]), i[1], i[2]))
                BtnSave['state'] = 'normal'
            else:
                messagebox.showinfo("Info", f"No Pending parts found for Quotation No. {n1}", parent=frm_delivery)
                
    def Swipe_To_Right():
        """Move selected item from List_Total_TreeView → List_Delivery_TreeView with custom qty."""
        selected = List_Total_TreeView.selection()
        if not selected:
            messagebox.showwarning("No selection", "Please select an item on the left.", parent=frm_delivery)
            return

        iid = selected[0]
        values = List_Total_TreeView.item(iid, "values")
        #print(values)
        if not values:
            return
        no, name, qty_str = values
        try:
            available_qty = int(qty_str)
        except ValueError:
            messagebox.showerror("Error", "Quantity must be a number.", parent=frm_delivery)
            return

        # Show custom quantity dialog
        dialog = QtyDialog(tab1, name, available_qty)
        tab1.wait_window(dialog)   # wait until dialog is closed

        qty_to_deliver = dialog.qty
        if qty_to_deliver is None:
            # User cancelled
            return

        # Update left tree with remaining qty
        remaining = available_qty - qty_to_deliver
        if remaining > 0:
            List_Total_TreeView.item(iid, values=(no, name, remaining))
        else:
            List_Total_TreeView.delete(iid)

        # Insert / merge into right tree
        merged = False
        for dest_iid in List_Delivery_TreeView.get_children():
            total_items = List_Delivery_TreeView.item(dest_iid, "values")
            dest_no, dest_name, dest_qty_str = total_items
            if dest_no == no and dest_name == name:
                try:
                    dest_qty = int(dest_qty_str)
                except ValueError:
                    dest_qty = 0
                List_Delivery_TreeView.item(dest_iid, values=(no, name, dest_qty + qty_to_deliver))
                merged = True
                break

        if not merged:
            List_Delivery_TreeView.insert("", "end", values=(no, name, qty_to_deliver))


    def Swipe_To_Left():
        """Move full quantity back from List_Delivery_TreeView → List_Total_TreeView."""
        selected = List_Delivery_TreeView.selection()
        if not selected:
            messagebox.showwarning("No selection", "Please select an item on the right.", parent=frm_delivery)
            return

        for iid in selected:
            values = List_Delivery_TreeView.item(iid, "values")
            if not values:
                continue
            no, name, qty_str = values
            try:
                qty = int(qty_str)
            except ValueError:
                qty = 0

            # Merge with left tree if same part found
            merged = False
            for src_iid in List_Total_TreeView.get_children():
                src_no, src_name, src_qty_str = List_Total_TreeView.item(src_iid, "values")
                if src_no == no and src_name == name:
                    try:
                        src_qty = int(src_qty_str)
                    except ValueError:
                        src_qty = 0
                    List_Total_TreeView.item(src_iid, values=(no, name, src_qty + qty))
                    merged = True
                    break

            if not merged:
                List_Total_TreeView.insert("", "end", values=(no, name, qty))

            List_Delivery_TreeView.delete(iid)

    def Generate_Excel_Report():
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        invoice_no = TxtInvoiceReport.get()
        sql1 = f"SELECT po_no FROM invoice_master where invoice_id = '{invoice_no}'"
        db_cursor.execute(sql1)
        data1 = db_cursor.fetchone()

        if not List_Total_TreeView_report.get_children():
            messagebox.showerror("Error", "No data found", parent=frm_delivery)
            return

        file_path = r'D:\\ToolCosting\\Support Documents\\Delivery_Report.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = "Delivery Report"

        # ----------- Collect Data Grouped by Challan -----------
        challan_data = defaultdict(list)

        for item_id in List_Total_TreeView_report.get_children():
            row = List_Total_TreeView_report.item(item_id, "values")

            challan_no = row[0]      # Adjust index if needed
            part_no = row[1]
            part_name = row[2]
            total_qty = row[3]
            delivered_qty = row[4]
            balance_qty = row[5]
            date = row[6]

            challan_data[challan_no].append(
                [part_no, part_name, total_qty, delivered_qty, balance_qty, date]
            )

        current_row = 1

        # ----------- Generate Report Per Challan -----------
        for challan_no, items in challan_data.items():

            # Title
            ws.merge_cells(start_row=current_row, start_column=1,
                        end_row=current_row, end_column=7)
            ws.cell(row=current_row, column=1).value = "CHALLAN REPORT"
            ws.cell(row=current_row, column=1).font = Font(size=14, bold=True)
            ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="center")

            current_row += 1

            # Invoice / PO / Challan header
            ws.cell(row=current_row, column=1).value = f"Invoice no. - {invoice_no}"
            ws.cell(row=current_row, column=3).value = f"PO No. - {data1[0]}"
            ws.cell(row=current_row, column=6).value = f"Challan No. - {str(challan_no).zfill(2)}"

            current_row += 2

            # Column headers
            headers = ["Sr. No.", "Part no.", "Part description",
                    "Total QTY", "Delivered QTY", "Balance QTY", "Date"]

            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=current_row, column=col_num)
                cell.value = header
                cell.font = Font(bold=True)
                cell.border = thin_border

            current_row += 1

            # Data rows
            
            
            sr_no = 1
           
            for item in items:
                row_data = [sr_no, item[0], item[1], item[2], item[3], item[4], item[5]]
                
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.value = value
                    cell.border = thin_border
                
                current_row += 1
                sr_no += 1

            # Space before next challan
            current_row += 3

        # ----------- Save File -----------
        try:
            wb.save(file_path)
            messagebox.showinfo("Success", f"Report saved to {file_path}", parent=frm_delivery)
            os.startfile(file_path)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save report: {e}", parent=frm_delivery)

    frm_delivery = Toplevel(master)
    frm_delivery.title("Challan Report")
    frm_delivery.geometry("950x520+300+170")
    frm_delivery.resizable(False, False)
    notebook = ttk.Notebook(frm_delivery)
    notebook.pack(expand=True, fill="both")

    tab1 = ttk.Frame(notebook)
    tab2 = ttk.Frame(notebook)

    notebook.add(tab1, text="                  Challan               ")
    notebook.add(tab2, text="                  Report                ")

    Lblhead = Label(tab1, text='New Challan',
                    font=('Times New Roman', 24, 'bold'))
    Lblhead.place(x=355, y=10)

    Frm1 = LabelFrame(tab1, width=430, height=60)
    Frm1.place(x=245, y=60)

    LblQuot = Label(tab1, text='Invoice No.', font=('Times New Roman', 16))
    LblQuot.place(x=255, y=75)

    TxtQuot = Entry(tab1, width=12, font=('Times New Roman', 16), justify='center')
    TxtQuot.place(x=375, y=75)
    TxtQuot.focus()

    BtnSearch = Button(tab1, text='Search',
                    font=('Times New Roman', 12),
                    width=12, bg='green', fg='white', command=Search_Records)
    BtnSearch.place(x=535, y=70)

    Lbl_Left_Tree = Label(tab1, text='Pending Delivery', font=('Times New Roman', 16, 'bold'), fg='brown')
    Lbl_Left_Tree.place(x=20, y=100)

    Lbl_Right_Tree = Label(tab1, text='Completed Delivery', font=('Times New Roman', 16, 'bold'), fg='brown')
    Lbl_Right_Tree.place(x=710, y=100)

    List_Total_TreeView = ttk.Treeview(
        tab1,
        columns=("no", "name", "qty")
    )

    List_Total_TreeView.heading("no", text="No.")
    List_Total_TreeView.heading("name", text="Part Name")
    List_Total_TreeView.heading("qty", text="Qty")

    List_Total_TreeView["show"] = "headings"

    List_Total_TreeView.column("no", width=70, anchor='center')
    List_Total_TreeView.column("name", width=150, anchor='center')
    List_Total_TreeView.column("qty", width=100, anchor='center')

    List_Total_TreeView.place(x=20, y=140, width=400, height=300)

    # Buttons between trees
    BtnRightArrow = Button(tab1, text='==>',
                        font=('Times New Roman', 12, 'bold'),
                        fg='white', bg='green', width=7,
                        command=Swipe_To_Right)
    BtnRightArrow.place(x=430, y=240)

    BtnLeftArrow = Button(tab1, text='<==',
                        font=('Times New Roman', 12, 'bold'),
                        fg='white', bg='green', width=7,
                        command=Swipe_To_Left)
    BtnLeftArrow.place(x=430, y=300)

    # Right Treeview - Delivery List
    scrolly2 = Scrollbar(tab1, orient=VERTICAL)
    scrolly2.place(x=920, y=140, height=300)

    List_Delivery_TreeView = ttk.Treeview(
        tab1,
        columns=("no", "name", "qty"),
        yscrollcommand=scrolly2.set
    )
    scrolly2.config(command=List_Delivery_TreeView.yview)

    List_Delivery_TreeView.heading("no", text="No.")
    List_Delivery_TreeView.heading("name", text="Part Name")
    List_Delivery_TreeView.heading("qty", text="Qty")

    List_Delivery_TreeView["show"] = "headings"

    List_Delivery_TreeView.column("no", width=70, anchor='center')
    List_Delivery_TreeView.column("name", width=150, anchor='center')
    List_Delivery_TreeView.column("qty", width=100, anchor='center')

    List_Delivery_TreeView.place(x=520, y=140, width=400, height=300)

    BtnSave = Button(tab1, text='Save', font=('Times New Roman', 14), width=14, bg='green', fg='white', command=Save_Records)
    BtnSave.place(x=400, y=450)

    On_Start()

    def f1(event):
        Search_Records()

    TxtQuot.bind('<Return>', f1)



    #=======================================================================================================
    #=======================================================================================================

    Lblhead = Label(tab2, text='Challan Report',
                    font=('Times New Roman', 24, 'bold'))
    Lblhead.place(x=335, y=10)


    Frm1_Report = LabelFrame(tab2, width=430, height=60)
    Frm1_Report.place(x=245, y=60)

    LblInvoiceReport = Label(tab2, text='Invoice No.', font=('Times New Roman', 16))
    LblInvoiceReport.place(x=255, y=75)

    TxtInvoiceReport = Entry(tab2, width=12, font=('Times New Roman', 16), justify='center')
    TxtInvoiceReport.place(x=375, y=75)

    BtnSearchReport = Button(tab2, text='Search',
                    font=('Times New Roman', 12),
                    width=12, bg='green', fg='white', command=Search_Record_Report)
    BtnSearchReport.place(x=535, y=70)
    
    BtnReport = Button(tab2, text='Generate Report', font=('Times New Roman', 14), width=17, bg='blue', fg='white', command=Generate_Excel_Report)
    BtnReport.place(x=380, y=450)

    List_Total_TreeView_report = ttk.Treeview(
        tab2,
        columns=("challan_no", "no", "name", "qty", "qty_del", "qty_pend", "dt")
    )

    List_Total_TreeView_report.heading("challan_no", text="Challan No.")
    List_Total_TreeView_report.heading("no", text="No.")
    List_Total_TreeView_report.heading("name", text="Part Name")
    List_Total_TreeView_report.heading("qty", text="Total Qty")
    List_Total_TreeView_report.heading("qty_del", text="Delivered Qty")
    List_Total_TreeView_report.heading("qty_pend", text="Remaining Qty")
    List_Total_TreeView_report.heading("dt", text="Date")

    List_Total_TreeView_report["show"] = "headings"

    List_Total_TreeView_report.column("challan_no", width=70, anchor='center')
    List_Total_TreeView_report.column("no", width=70, anchor='center')
    List_Total_TreeView_report.column("name", width=200, anchor='center')
    List_Total_TreeView_report.column("qty", width=100, anchor='center')
    List_Total_TreeView_report.column("qty_del", width=100, anchor='center')
    List_Total_TreeView_report.column("qty_pend", width=100, anchor='center')
    List_Total_TreeView_report.column("dt", width=100, anchor='center')

    List_Total_TreeView_report.place(x=20, y=140, width=910, height=300)

    def f2(evevnt):
        Search_Record_Report()

    TxtInvoiceReport.bind('<Return>', f2)
    frm_delivery.bind("<Escape>", lambda e: frm_delivery.destroy())
    # frm_delivery.mainloop()
    return frm_delivery
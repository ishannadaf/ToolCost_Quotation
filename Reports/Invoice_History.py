from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
import os
from tkinter import messagebox
from tkinter import *
from tkcalendar import DateEntry
from datetime import datetime
from Database.connection import *

def Frm_Invoice_Report(master, login_id):
    def Generate_Report(data, path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoice Report"
        # Header row
        headers = [
            "Sr. No.", "Invoice No", "GST No", "Customer Name", "Provider", "Amount", "Total GST", "PO No.", "Date"]
        ws.append(headers)

        # Data rows
        for row in data:
            ws.append(row)

        # ---------- STYLING ----------
        # Bold header, center alignment, borders
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if cell.row == 1:
                    cell.font = Font(bold=True)

        # Auto column width
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                value = str(cell.value) if cell.value is not None else ""
                max_length = max(max_length, len(value))
            ws.column_dimensions[col_letter].width = max_length + 2

        try:
            wb.save(path)
        except:
            messagebox.showerror("Error", "If you already opened report excel please close and retry.", parent=frm_report)
            return
        os.startfile(path)

    def Fetch_Data():
        date1 = datetime.strptime(TxtFrom.get(), '%d-%m-%Y').strftime('%Y-%m-%d')
        date2 = datetime.strptime(TxtTo.get(), '%d-%m-%Y').strftime('%Y-%m-%d')

        sql1 = f"""
        select im.invoice_id, am.gst_no, am.party_name, qm.provider_name, sum(qmd.total_price) as Amt, (sum(qmd.total_price) * 18 / 100) as Total_GST, im.po_no, qm.date_tr_quot from quotation_master qm 
        inner join account_master am on qm.cust_id = am.id
        inner join quotation_master_details qmd on qmd.Quotation_Id = qm.quotation_id
        inner join invoice_master im on im.quot_id = qm.quotation_id
        where im.date_tr between '{date1}' and '{date2}' 
        group by im.invoice_id, am.party_name, qm.provider_name, qm.date_tr_quot, am.gst_no, im.po_no
        """
        db_cursor.execute(sql1)
        data1 = db_cursor.fetchall()
        if data1:
            data = []
            cnt1 = 1
            for i in data1:
                data.append([cnt1, i[0], i[1], i[2], i[3], i[4], i[5], i[6], datetime.strftime(i[7], '%Y-%m-%d')])
                cnt1 += 1
            path = r"D:\\ToolCosting\\Support Documents\\Invoice_List.xlsx"
            messagebox.showinfo("Info", "Report Generate Successfully. Click Ok to open.", parent=frm_report)
            Generate_Report(data, path)
        else:
            messagebox.showerror("Error", f"No Invoice generated between {data1}, {date2}", parent=frm_report)
            

    frm_report = Toplevel(master)
    frm_report.geometry("500x300+480+220")
    frm_report.title("Invoice Report")
    frm_report.resizable(False, False)

    LblHead = Label(frm_report, text='Invoice Report', font=('Times New Roman', 24, 'bold'), fg='purple')
    LblHead.place(x=180, y=10)

    LblFrom = Label(frm_report, text='From', font=('Times New Roman', 18))
    LblFrom.place(x=90, y=80)
    TxtFrom = DateEntry(frm_report, width=12, font=('Times New Roman', 18), borderwidth=2, date_pattern="dd-mm-yyyy")
    TxtFrom.place(x=190, y=80)
    TxtFrom.focus_set()

    LblTo = Label(frm_report, text='To', font=('Times New Roman', 18))
    LblTo.place(x=90, y=150)
    TxtTo = DateEntry(frm_report, width=12, font=('Times New Roman', 18), borderwidth=2, date_pattern="dd-mm-yyyy")
    TxtTo.place(x=190, y=150)

    BtnGenerate = Button(frm_report, text='Generate', font=('Times New Roman', 14), width=14, bg='green', fg='white', command=Fetch_Data)
    BtnGenerate.place(x=190, y=210)
    frm_report.bind("<Escape>", lambda e: frm_report.destroy())
    # frm_report.mainloop()
    return frm_report
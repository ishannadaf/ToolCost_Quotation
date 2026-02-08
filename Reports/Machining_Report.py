from tkinter import *
from tkinter import messagebox
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font
from Database.connection import *
import os

def Frm_Machining_Report(master, login_id):
    def generate_quotation(items, filename, quote_id):
        operations = []
        seen = set()
        for item in items:
            for op in item["operations"].keys():
                if op not in seen:
                    seen.add(op)
                    operations.append(op)

        wb = Workbook()
        ws = wb.active
        ws.title = "Quotation"

        center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        header_font = Font(bold=True)

        # -------- Title (Quote Id) --------
        total_columns = 3 + 1 + len(operations) + 2   # SrNo,PartNo,Desc,RMC,ops,Profit,TOTAL
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
        c = ws.cell(row=1, column=1)
        c.value = f"Quot Id = {quote_id}"
        c.alignment = center
        c.font = header_font

        # Column widths (rough defaults)
        widths = [8, 10, 25] + [10] * (total_columns - 3)
        for col_idx, w in enumerate(widths, start=1):
            col_letter = chr(ord('A') + col_idx - 1)
            ws.column_dimensions[col_letter].width = w

        # -------- Helper: write one item block --------
        def write_item_block(start_row, item):
            # Header row: Price distribution (Rs)
            ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
            ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row,
                        end_column=total_columns)

            c = ws.cell(row=start_row, column=4)
            c.value = "Price distribution (Rs)"
            c.alignment = center
            c.font = header_font

            # Second header row (column titles)
            titles = ["Sr. No", "Part no", "Part description", "RMC"] \
                    + operations + ["Profit", "TOTAL"]

            for col, title in enumerate(titles, start=1):
                cell = ws.cell(row=start_row + 1, column=col, value=title)
                cell.alignment = center
                cell.font = header_font
                cell.border = thin_border

            # Data row
            data_row = start_row + 2

            # Build row values dynamically
            row_values = [
                item["sr_no"],
                item["part_no"],
                item["desc"],
                item["rmc"],
            ]

            # Add operation values in the same order as 'operations' list
            op_values = []
            for op in operations:
                amt = item["operations"].get(op, 0)
                op_values.append(amt)  # default 0 if not present

            profit = item["profit"]
            total = item["rmc"] + sum(op_values) + profit

            row_values.extend(op_values)
            row_values.append(profit)
            row_values.append(total)

            for col, val in enumerate(row_values, start=1):
                cell = ws.cell(row=data_row, column=col, value=val)
                if col == 3:  # description
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = center
                cell.border = thin_border

            # borders for merged header row
            for col in range(1, total_columns + 1):
                cell = ws.cell(row=start_row, column=col)
                cell.border = thin_border

            # return row index after leaving one blank row
            return data_row + 2

        # -------- Write all items --------
        current_row = 3
        for item in items:
            current_row = write_item_block(current_row, item)

        wb.save(filename)
        messagebox.showinfo("Info", "Report generated successfully. Click Ok to continue.", parent=frm_report)
        os.startfile(filename)
    def check_excel_open(filepath):
        try:
            # Try opening the file in append mode (fails if Excel has locked it)
            f = open(filepath, "a")
            f.close()
            return True
        except PermissionError:
            return False
        
    def Frm_Machining_Report_Excel(n1, path):
        sql1 = f"SELECT part_no, part_desc, rmc, unit_price, profit_per FROM quotation_master_details WHERE Quotation_Id = '{n1}'"
        db_cursor.execute(sql1)
        data1 = db_cursor.fetchall()
        items = []
        cnt = 1
        if data1:
            for i in data1:
                sql2 = f"SELECT machining_name, cust_rate, total_hr FROM quotation_master_details_machining_details WHERE '{n1}' AND part_no_id = '{i[0]}'"
                db_cursor.execute(sql2)
                data2 = db_cursor.fetchall()
                
                machining = {}
                total_sum = 0
                for j in data2:
                    machining[j[0]] = round(float(j[1]) * float(j[2]), 2)
                    total_sum += round(float(j[1]) * float(j[2]), 2)
                total_sum += i[2]
                profit_amt = round(total_sum * i[4] / 100, 2) 
                items.append({
                    "sr_no" : cnt,
                    "part_no" : i[0],
                    "desc": i[1],
                    "rmc": i[2],
                    "operations": machining,
                    "profit": profit_amt
                })
                cnt += 1

            generate_quotation(items, path, quote_id=n1)
        else:
            messagebox.showerror("Error", f"No record found for given Quotation ID : {n1}")
        
    def Generate_Report():
        n1 = TxtQuot.get()
        path = r"D:\\ToolCosting\\Support Documents\\Machining_Report.xlsx"
        print(n1)
        print(check_excel_open(path))
        if n1:
            if check_excel_open(path):
                Frm_Machining_Report_Excel(n1, path)
            else:
                messagebox.showerror("Error", "Excel is already opened. please close the excel.", parent=frm_report)

    frm_report = Toplevel(master)
    frm_report.geometry("500x220+500+250")
    frm_report.title("Machining Report")
    frm_report.resizable(False, False)
    
    LblHead = Label(frm_report, text='Machining Report', font=('Times New Roman', 22, 'bold'), fg='purple')
    LblHead.pack()
    
    Frm1 = LabelFrame(frm_report, width=400, height=70)
    Frm1.place(x=50, y=65)

    LblQuot = Label(frm_report, text='Quotation ID ', font=('Times New Roman', 18))
    LblQuot.place(x=90, y=80)
    TxtQuot = Entry(frm_report, font=('Times New Roman', 18), width=14, justify='center')
    TxtQuot.place(x=240, y=80)
    TxtQuot.focus()
    
    BtnGenerate = Button(frm_report, text='Generate', font=('Times New Roman', 14), width=14, bg='green', fg='white', command=Generate_Report)
    BtnGenerate.place(x=90, y=150)
    BtnExit = Button(frm_report, text='Exit', font=('Times New Roman', 14), width=14, bg='red', fg='white', command=frm_report.destroy)
    BtnExit.place(x=260, y=150)
    frm_report.bind("<Escape>", lambda e: frm_report.destroy())
    # frm_report.mainloop()
    return frm_report
    
# Frm_Machining_Report()